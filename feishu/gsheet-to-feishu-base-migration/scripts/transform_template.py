#!/usr/bin/env python3
"""清洗 + 去重 + 生成 diff.json + 台账 xlsx 的模板。每个项目复制一份到项目目录再改，别直接改这里。

输入:  gsheet_xlsx_to_json.py 产出的 sheet JSON；fbase.js schema/dump 产出的 schema_*.json / dump_*.json
输出:  diff.json（给 fbase.js apply）+ 迁移台账.xlsx（给人看：明细 / 丢弃行 / 合并去重 / 未匹配 / 映射口径）

要改的只有标 ★ 的地方。
"""
import json, re, datetime, collections, openpyxl
from openpyxl.styles import Font

# ★ 1. 目标表与源 ------------------------------------------------------------
APP = 'XXXX'                       # 多维表格 app_token（wiki 链接先 fbase.js resolve）
TBL = 'tblXXXX'                    # 目标数据表
SCHEMA = json.load(open(f'schema_{APP}.json'))
SOURCES = [                        # (标签, sheet json 路径, 表头所在行号(0起))
    ('源表A/某sheet', 'sheets/00_xxx.json', 0),
]
CURRENT = None                     # 增量迁移时填 dump_<tbl>.json；首次灌空表填 None

# ★ 2. 映射（全部要用户确认过，别自己猜） ---------------------------------------
PRODUCT = {}                       # 源写法(小写) → 飞书单选选项原文。映射不上 = 丢弃并进台账
PERSON = {}                        # 源人名 → open_id。拿 open_id: fbase.js people / persons
TIER_BY_FANS = [(10_000, 'Nano'), (100_000, 'Micro'), (500_000, 'Macro'), (1_000_000, 'BIG'), (float('inf'), 'Mega')]  # 量级按粉丝数重算，不信旧标签
COL = {                            # 飞书字段名 → 源表头（同一字段在不同源表头不同就写 list）
    # 'Influencer ID': ['Account', 'TikTok ID'],
}

# ---- 通用清洗函数（KANS 迁移踩出来的，直接用） ---------------------------------
def s(v): return '' if v is None else str(v).strip()
def handle(v):
    """达人ID 归一化：去@、去主页前缀、小写。唯一键用它，不用姓名（越南名拼写不统一）"""
    t = s(v); m = re.search(r'tiktok\.com/@([^/?\s]+)', t)
    return (m[1] if m else t).lstrip('@').strip().lower()
def num(v):
    """费用/数字：去货币符号和杂字符；1.2K/3M；越南写法 1.000.000 → 1000000"""
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    t = s(v).upper().replace('VND', '').replace('₫', '').replace('Đ', '').replace(' ', '')
    mult = 1
    if t.endswith('K'): mult, t = 1e3, t[:-1]
    elif t.endswith('M'): mult, t = 1e6, t[:-1]
    if re.fullmatch(r'\d{1,3}(\.\d{3})+', t): t = t.replace('.', '')       # 1.000.000
    t = t.replace(',', '')
    try: return float(t) * mult
    except ValueError: return None                                        # 非数字原值进备注，别丢
def pct(v):
    """百分比：'8,5%' 逗号当小数点；0.08 原样；'8%' → 0.08"""
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v) if v <= 1 else float(v) / 100
    t = s(v).replace(',', '.').replace(' ', '')
    m = re.fullmatch(r'(-?\d+(\.\d+)?)%?', t)
    if not m: return None
    x = float(m[1]); return x / 100 if (t.endswith('%') or x > 1) else x
def date(v):
    """→ datetime.date。ISO / dd/mm/yyyy / dd/mm（缺年按当年）。Google 会把 dd/mm 误转成 mm/dd，有月份线索时要纠正"""
    if v is None or v == '': return None
    t = s(v)
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', t)
    if m:
        try: return datetime.date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError: return None
    m = re.match(r'^(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?$', t)
    if m:
        y = int(m[3]) if m[3] else datetime.date.today().year
        if y < 100: y += 2000
        try: return datetime.date(y, int(m[2]), int(m[1]))
        except ValueError: return None
    return None
def ms(d): return int(datetime.datetime(d.year, d.month, d.day, tzinfo=datetime.timezone.utc).timestamp() * 1000) if d else None
def month_opt(d): return f'{d.month:02d}/{d.year}' if d else None
def video_id(v):
    m = re.search(r'/video/(\d{15,})', s(v)) or re.fullmatch(r'(\d{17,20})', s(v)); return m[1] if m else None
def vid_date(vid):
    """TikTok 视频ID>>32 = 发布时间戳（UTC，越南按 +7）。缺发布日时兜底"""
    return (datetime.datetime.fromtimestamp(int(vid) >> 32, datetime.timezone.utc) + datetime.timedelta(hours=7)).date() if vid else None
def tier(fans):
    f = num(fans)
    if f is None: return None
    return next(name for cap, name in TIER_BY_FANS if f < cap)
def person(name):
    oid = PERSON.get(s(name)); return [{'id': oid}] if oid else None
def link(u): return {'link': u, 'text': u} if s(u).startswith('http') else None

# 字段写入格式：按 schema 类型自动转换。公式/查找引用/自动编号/系统字段一律不写
SKIP_TYPES = {'公式', '查找引用', '自动编号', '创建时间', '修改时间', '创建人', '修改人', '双向关联', '单向关联'}
FIELDS = {f['name']: f for f in SCHEMA[TBL]['fields']}
def put(rec, name, val, unmatched):
    f = FIELDS.get(name)
    if f is None: raise KeyError(f'飞书没有字段「{name}」——字段可能被团队改名了，重跑 schema')
    if val in (None, '', []) or f['type'] in SKIP_TYPES: return
    if f['type'] in ('单选', '多选'):
        opts = {o['name'] for o in f['options'] or []}
        vals = val if isinstance(val, list) else [val]
        ok = [v for v in vals if v in opts]
        for v in vals:
            if v not in opts: unmatched[(name, v)] += 1     # 选项不自己加，留空并记台账
        if not ok: return
        val = ok if f['type'] == '多选' else ok[0]
    rec[name] = val

# ---- 读源 → 规范化 ---------------------------------------------------------
def read_sources():
    rows = []
    for label, path, hrow in SOURCES:
        data = json.load(open(path)); H = [s(h).replace('\n', ' ') for h in data[hrow]]
        idx = lambda names: next((H.index(n) for n in ([names] if isinstance(names, str) else names) if n in H), None)
        for i, r in enumerate(data[hrow + 1:], start=hrow + 2):
            r = list(r) + [None] * len(H)
            g = lambda name: r[idx(COL[name])] if idx(COL[name]) is not None else None
            rows.append(dict(_src=label, _row=i, _get=g))
    return rows

def build():
    unmatched = collections.Counter(); dropped = []; kept = []
    for x in read_sources():
        g = x['_get']
        # ★ 3. 一行源数据 → 飞书字段。示例：
        acc = handle(g('Influencer ID')) if 'Influencer ID' in COL else ''
        prod = PRODUCT.get(s(g('Product')).lower()) if 'Product' in COL else None
        if not acc: dropped.append((x['_src'], x['_row'], '', '', '达人ID为空')); continue
        if 'Product' in COL and not prod: dropped.append((x['_src'], x['_row'], acc, s(g('Product')), '产品为空或无对应选项')); continue
        rec, note = {}, []
        put(rec, 'Influencer ID', acc, unmatched)
        if prod: put(rec, 'Product', prod, unmatched)
        # put(rec, 'BD', person(g('BD')), unmatched); if g('BD') and not person(g('BD')): note.append('原表BD: '+s(g('BD')))
        # put(rec, 'Fee VND', num(g('Fee VND')), unmatched); if g('Fee VND') and num(g('Fee VND')) is None: note.append('【Fee】'+s(g('Fee VND')))
        # put(rec, 'Update Date', ms(date(g('Update Date'))), unmatched)
        # put(rec, 'Video Link', link(g('Video Link')), unmatched)
        # if note: put(rec, 'Note', ' | '.join(note), unmatched)
        x['rec'] = rec; x['key'] = (acc, prod)          # ★ 4. 去重键。有报价：(达人, 报价, 产品, BD)；有视频：Video ID 最安全
        kept.append(x)

    # 去重：同键多行合并成一行（★ 5. 选哪行为准：进度最靠后 > 日期最新）
    groups = collections.defaultdict(list)
    for x in kept: groups[x['key']].append(x)
    final, merged = [], []
    for k, g in groups.items():
        best = g[-1]
        if len(g) > 1: merged.append((*map(str, k), len(g), '; '.join(f"{y['_src']}#{y['_row']}" for y in g)))
        final.append(best)

    # 增量：和飞书现状比对，只新增键不存在的
    diff = {'create': [], 'update': [], 'delete': []}
    have = set()
    if CURRENT:
        for it in json.load(open(CURRENT)):
            f = it['fields']
            have.add((handle(f.get('Influencer ID', [{}])[0].get('text') if isinstance(f.get('Influencer ID'), list) else f.get('Influencer ID')), f.get('Product')))
    for x in final:
        if x['key'] not in have: diff['create'].append(x['rec'])
    json.dump({TBL: diff}, open('diff.json', 'w'), ensure_ascii=False)

    # 台账
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '待写入'
    cols = sorted({k for x in final for k in x['rec']}); ws.append(cols + ['来源', '源行'])
    for x in final:
        ws.append([json.dumps(x['rec'].get(c), ensure_ascii=False) if isinstance(x['rec'].get(c), (dict, list)) else x['rec'].get(c) for c in cols] + [x['_src'], x['_row']])
    w = wb.create_sheet('丢弃行'); w.append(['来源', '源行', '达人', '原值', '原因']); [w.append(list(d)) for d in dropped]
    w = wb.create_sheet('合并去重'); w.append(['键…', '', '源行数', '来源行']); [w.append(list(m)) for m in merged]
    w = wb.create_sheet('选项未匹配'); w.append(['字段', '原值', '行数']); [w.append([k[0], k[1], n]) for k, n in unmatched.most_common()]
    w = wb.create_sheet('映射口径'); w.append(['产品映射', json.dumps(PRODUCT, ensure_ascii=False)]); w.append(['人员映射', ', '.join(PERSON)])
    for sh in wb.worksheets:
        for c in sh[1]: c.font = Font(bold=True)
    wb.save('迁移台账.xlsx')
    print(f'源行 {len(kept)+len(dropped)}  保留 {len(kept)}  丢弃 {len(dropped)}  去重后 {len(final)}  新增 {len(diff["create"])}')
    print('丢弃原因', collections.Counter(d[4] for d in dropped).most_common())
    print('选项未匹配 Top', unmatched.most_common(15))

if __name__ == '__main__':
    build()
