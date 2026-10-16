#!/usr/bin/env python3
"""KANS 越南｜C-ads 品广素材投放决策汇总 —— 从平台导出文件生成汇总表 xlsx。

用法：
  python3 build_summary.py --fx 26310 --out 汇总.xlsx  文件或文件夹 [文件 ...]

自动识别的输入（按内容识别，不靠文件名）：
  1. TTMS 种草广告报告：报告页「全选复制」粘贴成的 .txt（可多份拼在一个 txt 里），
     或按 templates/TTMS结案录入模板.csv 手填的 csv/xlsx            —— 必需
  2. GMV MAX Creative data 导出 xlsx（含 Post ID / Cost / Gross revenue，可多份）—— 强烈建议
  3. 达人订单明细 affiliate_orders_*.csv（含 内容ID / 店铺广告佣金率）            —— 可选，算 V/W
  4. 广告后台 Campaign Report xlsx（含 广告组名称 / 一级状态 / 广告组预算）       —— 可选，取 KOL/状态/日预算
"""
import argparse
import glob
import math
import os
import re
import sys
from collections import Counter, OrderedDict

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- 常量 / 口径

PRODUCT_SHORT = OrderedDict([  # 报告名/广告组名里的产品写法 → 表里的简称
    ('white essence', '白'), ('white', '白'), ('白精华', '白'), ('白', '白'),
    ('red serum', '红'), ('red', '红'), ('红精华', '红'), ('红', '红'),
    ('tone-up cream', '素颜霜'), ('tone up cream', '素颜霜'), ('素颜霜', '素颜霜'),
    ('377次抛', '377次抛'), ('377 serum', '377次抛'), ('377', '377次抛'),
    ('战痘次抛', '战痘次抛'), ('acne serum', '战痘次抛'), ('acne', '战痘次抛'), ('战痘', '战痘次抛'),
])
GMVMAX_PRODUCT_ID = {  # GMV MAX 导出文件名里的 Product ID → 简称（示例 ID，换成你自己的商品 ID）
    '1000000000000000001': '白', '1000000000000000002': '红', '1000000000000000003': '素颜霜',
    '1000000000000000004': '377次抛', '1000000000000000005': '战痘次抛',
}
REPORT_NAME_RE = re.compile(r'(\d{18,20})_(\d{4})_(\d{4})_([^\n\r\t|]+)')

HEAD = ['序号', '素材(KOL)', '视频ID', '产品', '状态',
        '结案消耗($)', '结案CPCo($)', '结案种草Co', '6sVTR', '种草效率', '结案GMV($)', '结案ROAS', '品广质量分(结案)',
        'GMVMAX消耗($)', 'GMVMAX ROI', 'GMVMAX收入($)', 'GMVMAX获客成本($)',
        '总营销ROI(max收入/所有消耗)', '品广CPM($)', 'GMVMAX CPM($)', '总CPM($)', 'GMV付费流占比', 'GMV自然流占比',
        '现行日预算($)', '新预算($)', '投流码']
GROUPS = [(6, 13, '▶ 投后结案（TTMS 种草广告报告 · 主口径）', '2F5597'),
          (14, 17, 'GMV MAX（辅助对照 · 素材口径）', '1E7145'),
          (18, 26, '综合 / 决策', '2F5597')]

WARN = []  # 校验提示，最后写进「校验」页并打印


def warn(msg):
    WARN.append(msg)


def short_product(s):
    if s is None:
        return ''
    t = str(s).strip().lower()
    for k, v in PRODUCT_SHORT.items():
        if k in t:
            return v
    return str(s).strip()


def num(s):
    """'1,234.5' / '12.3%' / '¥6.4' / '$7' → float；'-' / 空 → None。百分号按原数返回（不除 100）。"""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return None if (isinstance(s, float) and math.isnan(s)) else float(s)
    t = str(s).strip().replace(',', '').replace('，', '')
    t = re.sub(r'^[¥$₫€]|USD$|VND$', '', t).strip().rstrip('%').strip()
    try:
        return float(t)
    except ValueError:
        return None


def pct(s):
    """'8.01%' → 0.0801；'0.0801' → 0.0801。"""
    if s is None:
        return None
    if isinstance(s, str) and '%' in s:
        v = num(s)
        return None if v is None else v / 100
    v = num(s)
    if v is None:
        return None
    return v / 100 if v > 1 else v


# ---------------------------------------------------------------- 1. TTMS 报告页粘贴文本

# 字段 → 可能出现的标签（英文界面为准；中文界面标签为兜底猜测，认不出时用录入模板）
TTMS_LABELS = {
    'imp':  [r'Impressions?', r'展示次数', r'展示量', r'曝光量', r'曝光次数', r'曝光数'],
    'cost': [r'Cost', r'消耗', r'花费', r'总消耗'],
    'cpm':  [r'CPM', r'千次展示成本', r'千次展示费用'],
    'vtr6': [r'6-second view-through rate', r'6\s*秒完播率', r'6\s*秒播放率', r'6\s*秒观看率'],
    'v6':   [r'6-second video views', r'6\s*秒播放次数', r'6\s*秒观看次数', r'6\s*秒播放量'],
    'co':   [r'New Consideration size', r'新增种草人群规模', r'新增种草人数', r'新增种草人群', r'新意向受众规模'],
    'cor':  [r'New Consideration rate', r'新增种草率', r'新增种草人群占比', r'新意向受众率'],
    'gmv':  [r'GMV', r'成交金额'],
    'roas': [r'ROAS', r'广告支出回报率'],
}
TTMS_PCT = {'vtr6', 'cor'}
VALUE_LINE = re.compile(r'^[¥$₫]?\s*-?[\d,]+(\.\d+)?\s*%?$')


def _find_value(lines, labels, want_pct):
    for i, ln in enumerate(lines):
        s = ln.strip()
        for lb in labels:
            m = re.match(r'^(?:%s)\s*[:：]?\s*(.*)$' % lb, s, flags=re.I)
            if not m:
                continue
            rest = re.sub(r'\s*(USD|VND)$', '', m.group(1).strip())
            if rest and not VALUE_LINE.match(rest):
                continue  # 「Cost per Consideration」这类更长的标签，不是要找的那个
            cand = [rest] if rest else [x.strip() for x in lines[i + 1:i + 3] if x.strip()][:1]
            for c in cand:
                c0 = re.sub(r'\s*(USD|VND)$', '', c)
                if VALUE_LINE.match(c0) and (('%' in c0) == want_pct or want_pct is None):
                    return c0
    return None


def parse_ttms_text(text, src):
    """一个 txt 可以粘多份报告。按报告名切块：同名连续出现算同一块。"""
    hits = [(m.start(), m.group(0).strip()) for m in REPORT_NAME_RE.finditer(text)]
    if not hits:
        warn(f'{src}：没找到形如「素材ID_起投MMDD_截止MMDD_产品」的报告名，这个文件没读进来')
        return []
    starts = []
    for pos, name in hits:
        if not starts or starts[-1][1] != name:
            starts.append((pos, name))
    rows = []
    for k, (pos, name) in enumerate(starts):
        end = starts[k + 1][0] if k + 1 < len(starts) else len(text)
        block = text[pos:end]
        lines = block.splitlines()
        r = {'report': name, 'src': os.path.basename(src)}
        cyc = re.search(r'(\d{4}-\d{2}-\d{2})\s*[~～至-]+\s*(\d{4}-\d{2}-\d{2})', block)
        r['cycle'] = f'{cyc.group(1)} ~ {cyc.group(2)}' if cyc else None
        for f, labels in TTMS_LABELS.items():
            raw = _find_value(lines, labels, f in TTMS_PCT)
            r[f + '_raw'] = raw
            r[f] = pct(raw) if f in TTMS_PCT else num(raw)
        rows.append(r)
    return rows


TEMPLATE_COLS = {  # 手填模板列名 → 字段
    'report': ['报告名', '报告名称', 'report'], 'vid': ['视频ID', '素材ID', 'video_id'],
    'product': ['产品'], 'cost': ['消耗(USD)', '消耗', '结案消耗'], 'imp': ['展示量', '曝光', '展示次数', 'Impression'],
    'co': ['种草Co', '新增种草人数', 'New Consideration size'], 'v6': ['6秒播放数', '6-second video views'],
    'vtr6': ['6sVTR', '6秒完播率'], 'cor': ['种草效率', 'New Consideration rate'],
    'gmv': ['GMV(USD)', 'GMV'], 'roas': ['ROAS'], 'cpm': ['CPM(USD)', 'CPM'], 'cycle': ['分析周期'],
}


def parse_ttms_table(df, src):
    col = {}
    for f, names in TEMPLATE_COLS.items():
        for n in names:
            for c in df.columns:
                if str(c).strip().lower() == n.lower():
                    col[f] = c
                    break
            if f in col:
                break
    rows = []
    for _, x in df.iterrows():
        r = {'src': os.path.basename(src)}
        rep = str(x[col['report']]).strip() if 'report' in col and pd.notna(x[col['report']]) else ''
        vid = str(x[col['vid']]).strip() if 'vid' in col and pd.notna(x[col['vid']]) else ''
        vid = re.sub(r'\.0$', '', vid)
        if not rep and not vid:
            continue
        if not rep:
            rep = f"{vid}_0000_0000_{x[col['product']] if 'product' in col else ''}"
        r['report'] = rep
        for f in ['cost', 'imp', 'co', 'v6', 'gmv', 'roas', 'cpm']:
            r[f] = num(x[col[f]]) if f in col else None
            r[f + '_raw'] = None if f not in col else x[col[f]]
        for f in ['vtr6', 'cor']:
            r[f] = pct(x[col[f]]) if f in col and pd.notna(x[col[f]]) else None
            r[f + '_raw'] = None if f not in col else x[col[f]]
        r['cycle'] = str(x[col['cycle']]) if 'cycle' in col and pd.notna(x[col['cycle']]) else None
        rows.append(r)
    return rows


def ttms_by_video(rows):
    """校验每份报告 → 按素材ID合并（同素材多份报告：量相加、比率重算）。"""
    seen = {}
    for r in rows:
        if r['report'] in seen:
            warn(f"报告「{r['report']}」粘了不止一次，只保留最后一份")
        seen[r['report']] = r
    per_vid = OrderedDict()
    for r in seen.values():
        m = REPORT_NAME_RE.match(r['report'])
        vid, start, end, prod = (m.group(1), m.group(2), m.group(3), m.group(4).strip()) if m else (r['report'], '', '', '')
        miss = [k for k in ['cost', 'imp', 'co'] if r.get(k) is None]
        if miss:
            warn(f"报告「{r['report']}」缺 {'/'.join(miss)}，这一份没进表（检查是不是没复制全）")
            continue
        chk = []
        if r.get('cor') is not None and r['imp']:
            c = r['co'] / r['imp']
            if abs(c - r['cor']) > max(0.03 * r['cor'], 0.00006):
                chk.append(f"种草效率 页面 {r['cor']:.2%} ≠ 种草Co÷展示量 {c:.2%}")
        if r.get('vtr6') is not None and r.get('v6') is not None and r['imp']:
            c = r['v6'] / r['imp']
            if abs(c - r['vtr6']) > max(0.03 * r['vtr6'], 0.00006):
                chk.append(f"6sVTR 页面 {r['vtr6']:.2%} ≠ 6秒播放÷展示量 {c:.2%}")
        if r.get('cpm') is not None and r['imp']:
            c = r['cost'] / r['imp'] * 1000
            if abs(c - r['cpm']) > max(0.03 * r['cpm'], 0.006):
                chk.append(f"CPM 页面 {r['cpm']:.2f} ≠ 消耗÷展示量×1000 {c:.2f}")
        if chk:
            warn(f"报告「{r['report']}」页面数字互相对不上（多半复制时页面没加载完，建议重新复制这一份）：" + '；'.join(chk))
        if r.get('v6') is None and r.get('vtr6') is not None:
            r['v6'] = r['vtr6'] * r['imp']
        if r.get('v6') is None:
            warn(f"报告「{r['report']}」缺 6 秒播放数据，6sVTR 留空")
        # GMV：优先页面值；页面缺失或和「消耗×ROAS」差 >2%（页面数字还在滚动动画时复制会这样）→ 用反算值
        exp = r['cost'] * r['roas'] if r.get('roas') is not None else None
        g = r.get('gmv')
        if g is None and exp is not None:
            r['gmv_final'], r['gmv_note'] = exp, '页面缺 GMV，已用 消耗×ROAS 反算'
        elif g is not None and exp is not None and exp > 0 and abs(g - exp) / exp > 0.02:
            r['gmv_final'], r['gmv_note'] = exp, f"页面 GMV={r.get('gmv_raw')} 与 消耗×ROAS={exp:,.2f} 对不上，已用反算值"
        else:
            r['gmv_final'], r['gmv_note'] = g, ''
        if r['gmv_note']:
            warn(f"报告「{r['report']}」：{r['gmv_note']}")
        d = per_vid.setdefault(vid, {'vid': vid, 'reports': [], 'prod': prod, 'cost': 0.0, 'imp': 0.0, 'co': 0.0,
                                     'v6': 0.0, 'v6_ok': True, 'gmv': 0.0, 'gmv_ok': True, 'cycles': set()})
        d['reports'].append(r['report'])
        d['prod'] = d['prod'] or prod
        for k in ['cost', 'imp', 'co']:
            d[k] += r[k]
        if r.get('v6') is None:
            d['v6_ok'] = False
        else:
            d['v6'] += r['v6']
        if r.get('gmv_final') is None:
            d['gmv_ok'] = False
        else:
            d['gmv'] += r['gmv_final']
        if r.get('cycle'):
            d['cycles'].add(r['cycle'])
        d['single'] = r if len(d['reports']) == 1 else None
    for d in per_vid.values():
        if len(d['reports']) > 1:
            warn(f"素材 {d['vid']} 有 {len(d['reports'])} 份报告，已合并：{'、'.join(d['reports'])}")
            d['single'] = None
    return per_vid


# ---------------------------------------------------------------- 2. GMV MAX

def load_gmvmax(paths, fx):
    frames, labels = [], []
    for p in paths:
        x = pd.read_excel(p, dtype={'Post ID': str})
        x = x[x['Post ID'].notna() & x['Post ID'].astype(str).str.fullmatch(r'\d{15,20}')]
        frames.append(x)
        base = os.path.basename(p)
        pid = re.search(r'(\d{19})', base)
        lab = GMVMAX_PRODUCT_ID.get(pid.group(1)) if pid else None
        if not lab:
            stem = os.path.splitext(base)[0].lower()
            lab = next((v for k, v in PRODUCT_SHORT.items() if len(k) > 3 and k in stem), None)
        labels.append(lab or os.path.splitext(base)[0])
    if not frames:
        return {}, []
    g = pd.concat(frames, ignore_index=True)
    for c in ['Cost', 'Gross revenue', 'SKU orders', 'Product ad impressions']:
        g[c] = pd.to_numeric(g[c], errors='coerce').fillna(0)
    cur = g['Currency'].astype(str).str.upper() if 'Currency' in g else pd.Series(['VND'] * len(g))
    div = cur.map(lambda c: 1.0 if c == 'USD' else fx)
    g['cost_usd'] = g['Cost'] / div
    g['rev_usd'] = g['Gross revenue'] / div
    agg = g.groupby('Post ID').agg(cost=('cost_usd', 'sum'), rev=('rev_usd', 'sum'), orders=('SKU orders', 'sum'),
                                   imp=('Product ad impressions', 'sum'), acct=('TikTok account', 'first'))
    return agg.to_dict('index'), list(OrderedDict.fromkeys(labels))


# ---------------------------------------------------------------- 3. 达人订单（付费流占比）

ORDER_COLS = {'cid': ['内容ID', 'Content ID'], 'rate': ['店铺广告佣金率', 'Shop ads commission rate'],
              'status': ['订单状态', 'Order status'], 'amt': ['支付金额', 'Payment amount', 'Payment Amount']}


def load_orders(paths):
    out = {}
    for p in paths:
        a = pd.read_csv(p, dtype=str, encoding='utf-8-sig')
        a.columns = [c.strip() for c in a.columns]
        col = {}
        for f, names in ORDER_COLS.items():
            col[f] = next((n for n in names if n in a.columns), None)
        if None in col.values():
            warn(f'{os.path.basename(p)}：订单表缺列 {[f for f, c in col.items() if c is None]}，V/W 没算')
            continue
        st = a[col['status']].astype(str).str.strip()
        a = a[st.isin(['已结算', 'Settled', 'settled', 'Completed'])].copy()
        a['amt'] = pd.to_numeric(a[col['amt']], errors='coerce').fillna(0)
        r = a[col['rate']].fillna('').astype(str).str.strip()
        a['paid'] = r.ne('') & r.ne('0%') & r.ne('0')
        for cid, s in a.groupby(col['cid']):
            o = out.setdefault(str(cid).strip(), {'tot': 0.0, 'paid': 0.0, 'n': 0})
            o['tot'] += s['amt'].sum()
            o['paid'] += s.loc[s['paid'], 'amt'].sum()
            o['n'] += len(s)
    return out


# ---------------------------------------------------------------- 4. Campaign Report（KOL / 状态 / 日预算）

CAMP_COLS = {'name': ['广告组名称', 'Ad group name'], 'status': ['一级状态', 'Primary status', 'Status'],
             'budget': ['广告组预算', 'Ad group budget', 'Budget']}


def load_campaign(paths):
    if not paths:
        return {}, None
    if len(paths) > 1:
        paths = sorted(paths, key=lambda p: re.findall(r'\d{4}-\d{2}-\d{2}', p)[-1:] or [''])
        warn(f'给了 {len(paths)} 份 Campaign Report，状态/日预算取日期最新的：{os.path.basename(paths[-1])}')
    p = paths[-1]
    c = pd.read_excel(p, dtype=str)
    col = {f: next((n for n in names if n in c.columns), None) for f, names in CAMP_COLS.items()}
    if col['name'] is None:
        warn(f'{os.path.basename(p)}：找不到「广告组名称」列，状态/日预算没取')
        return {}, None
    c = c[c[col['name']].notna() & ~c[col['name']].astype(str).str.contains('总计|Total', regex=True)]
    out = {}
    for _, x in c.iterrows():
        name = str(x[col['name']])
        vid = re.search(r'(\d{18,20})', name)
        if not vid:
            continue
        seg = [s.strip() for s in name.split('-')]
        d = out.setdefault(vid.group(1), {'kol': seg[1] if len(seg) > 1 else '', 'status': [], 'budget': 0.0,
                                          'groups': []})
        st = str(x[col['status']]).strip() if col['status'] and pd.notna(x[col['status']]) else ''
        bud = num(x[col['budget']]) if col['budget'] else None
        d['status'].append(st)
        d['groups'].append(name)
        if bud is not None and st in ('投放中', 'Active', 'Delivering', ''):
            d['budget'] += bud
    dates = re.findall(r'\d{4}-\d{2}-\d{2}', os.path.basename(p))
    return out, (dates[-1] if dates else None)


# ---------------------------------------------------------------- 文件分类

def classify(paths):
    kinds = {'ttms_text': [], 'ttms_table': [], 'gmvmax': [], 'orders': [], 'campaign': [], 'unknown': []}
    files = []
    for p in paths:
        if os.path.isdir(p):
            for q in sorted(glob.glob(os.path.join(p, '**', '*'), recursive=True)):
                if os.path.isfile(q):
                    files.append(q)
        elif os.path.isfile(p):
            files.append(p)
        else:
            warn(f'找不到文件：{p}')
    for f in files:
        base = os.path.basename(f)
        ext = os.path.splitext(f)[1].lower()
        if base.startswith(('.', '~$')):
            continue
        try:
            if ext in ('.txt', '.md'):
                kinds['ttms_text'].append(f)
            elif ext == '.csv':
                head = pd.read_csv(f, nrows=0, dtype=str, encoding='utf-8-sig').columns.str.strip().tolist()
                if any(c in head for c in ORDER_COLS['cid']):
                    kinds['orders'].append(f)
                elif any(c in head for c in TEMPLATE_COLS['report'] + TEMPLATE_COLS['vid']):
                    kinds['ttms_table'].append(f)
                else:
                    kinds['unknown'].append(f)
            elif ext in ('.xlsx', '.xls'):
                head = [str(c).strip() for c in pd.read_excel(f, nrows=0).columns]
                if 'Post ID' in head and ('Gross revenue' in head or 'Cost' in head):
                    kinds['gmvmax'].append(f)
                elif any(c in head for c in CAMP_COLS['name']):
                    kinds['campaign'].append(f)
                elif any(c in head for c in TEMPLATE_COLS['report'] + TEMPLATE_COLS['vid']):
                    kinds['ttms_table'].append(f)
                else:
                    kinds['unknown'].append(f)
            else:
                kinds['unknown'].append(f)
        except Exception as e:  # noqa: BLE001
            warn(f'{base} 读不了：{e}')
    return kinds


# ---------------------------------------------------------------- 计算

def mm(vals):
    v = [x for x in vals if x is not None]
    lo, hi = (min(v), max(v)) if v else (0, 0)
    return [None if x is None else (0.5 if hi == lo else (x - lo) / (hi - lo)) for x in vals]


def build_rows(vids, gm, orders, camp):
    rows = []
    for vid, d in vids.items():
        s = d.get('single')
        r = {'vid': vid, 'prod': short_product(d['prod']), 'F': d['cost'], 'H': d['co']}
        r['G'] = d['cost'] / d['co'] if d['co'] else None
        # 单份报告：比率直接用报告页数值（与平台一致）；多份合并：用合计重算
        if s is not None and s.get('vtr6') is not None:
            r['I'] = s['vtr6']
        else:
            r['I'] = d['v6'] / d['imp'] if d['v6_ok'] and d['imp'] else None
        r['J'] = s['cor'] if s is not None and s.get('cor') is not None else (d['co'] / d['imp'] if d['imp'] else None)
        r['K'] = d['gmv'] if d['gmv_ok'] else None
        if s is not None and s.get('roas') is not None:
            r['L'] = s['roas']
        else:
            r['L'] = r['K'] / d['cost'] if r['K'] is not None and d['cost'] else None
        r['S'] = d['cost'] / d['imp'] * 1000 if d['imp'] else None
        r['imp_c'] = d['imp']
        g = gm.get(vid)
        if g:
            r['N'], r['P'] = g['cost'], g['rev']
            r['O'] = g['rev'] / g['cost'] if g['cost'] else None
            r['Q'] = g['cost'] / g['orders'] if g['orders'] else None
            r['R'] = g['rev'] / (d['cost'] + g['cost']) if (d['cost'] + g['cost']) else None
            r['T'] = g['cost'] / g['imp'] * 1000 if g['imp'] >= 1000 else None
            r['U'] = (d['cost'] + g['cost']) / (d['imp'] + g['imp']) * 1000 if g['imp'] >= 1000 else None
            r['acct'] = g.get('acct')
        else:
            r['N'] = '无数据'
            warn(f'素材 {vid} 在 GMV MAX 导出里没找到（可能没进 GMV MAX 或少传了某个产品的文件），N–R/T/U 留空')
        o = orders.get(vid)
        if o and o['tot'] > 0:
            r['V'] = o['paid'] / o['tot']
            r['W'] = 1 - r['V']
            r['n_orders'] = o['n']
        c = camp.get(vid)
        if c:
            sts = [x for x in c['status'] if x]
            r['E'] = '投放中' if '投放中' in sts else (Counter(sts).most_common(1)[0][0] if sts else '')
            r['X'] = c['budget'] if c['budget'] else None
            r['B'] = c['kol']
        if not r.get('B'):
            r['B'] = r.get('acct') or ''
        r['reports'] = d['reports']
        rows.append(r)
    # 品广质量分：批内等权 min-max（6sVTR 越高越好、种草效率越高越好、CPCo 越低越好）
    a, b, c = mm([r['I'] for r in rows]), mm([r['J'] for r in rows]), mm([r['G'] for r in rows])
    for r, x, y, z in zip(rows, a, b, c):
        r['M'] = None if None in (x, y, z) else (x + y + (1 - z)) / 3
    # 序号 = 质量分排名；行顺序 = 总营销ROI 从高到低（无 GMV MAX 的排最后）
    for i, r in enumerate(sorted(rows, key=lambda r: -(r['M'] if r['M'] is not None else -1)), 1):
        r['A'] = i
    # R 保留两位后并列时，质量分高的排前面
    rows.sort(key=lambda r: (r.get('R') is None, -round(r.get('R') or 0, 2), -(r['M'] if r['M'] is not None else -1)))
    return rows


# ---------------------------------------------------------------- 写 xlsx

FMT = {'F': '#,##0.00', 'G': '0.0000', 'H': '#,##0', 'I': '0.00%', 'J': '0.00%', 'K': '#,##0.00', 'L': '0.00',
       'M': '0.000', 'N': '#,##0.00', 'O': '0.00', 'P': '#,##0.00', 'Q': '0.00', 'R': '0.00', 'S': '0.00',
       'T': '0.00', 'U': '0.00', 'V': '0.0%', 'W': '0.0%', 'X': '#,##0', 'Y': '#,##0'}
RND = {'F': 2, 'G': 4, 'I': 4, 'J': 4, 'K': 2, 'L': 2, 'M': 3, 'N': 2, 'O': 2, 'P': 2, 'Q': 2, 'R': 2, 'S': 2,
       'T': 2, 'U': 2, 'V': 3, 'W': 3}


def write_xlsx(rows, out, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = '汇总'
    F = '微软雅黑'
    thin = Side(style='thin', color='D0D7E5')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ncol = len(HEAD)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, 'KANS 越南｜C-ads 品广素材投放决策汇总')
    c.font = Font(name=F, size=14, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='2F5597')
    c.alignment = Alignment(horizontal='center', vertical='center')
    for rr, txt in ((2, meta['line2']), (3, meta['line3'])):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=ncol)
        c = ws.cell(rr, 1, txt)
        c.font = Font(name=F, size=9, color='404040')
        c.fill = PatternFill('solid', fgColor='EEF2F8')
        c.alignment = Alignment(horizontal='right' if rr == 2 else 'left', vertical='center', wrap_text=True)
    for c1, c2, title, color in GROUPS:
        ws.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
        c = ws.cell(4, c1, title)
        c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=color)
        c.alignment = Alignment(horizontal='center', vertical='center')
    for i, h in enumerate(HEAD, 1):
        c = ws.cell(5, i, h)
        c.font = Font(name=F, size=10, bold=True)
        c.fill = PatternFill('solid', fgColor='DCE6F1')
        c.border = bd
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    letters = [get_column_letter(i) for i in range(1, ncol + 1)]
    yel = PatternFill('solid', fgColor='FFF2CC')
    for k, r in enumerate(rows):
        R = 6 + k
        vals = {'A': r['A'], 'B': r.get('B', ''), 'C': r['vid'], 'D': r['prod'], 'E': r.get('E', ''),
                'Y': None, 'Z': None}
        for L in 'FGHIJKLMNOPQRSTUVWX':
            v = r.get(L)
            vals[L] = round(v, RND[L]) if isinstance(v, float) and L in RND else v
        for i, L in enumerate(letters, 1):
            c = ws.cell(R, i, vals.get(L))
            c.font = Font(name=F, size=10)
            c.border = bd
            c.alignment = Alignment(horizontal='left' if L in 'BC' else 'center', vertical='center')
            if L in FMT and isinstance(vals.get(L), (int, float)):
                c.number_format = FMT[L]
            if L == 'C':
                c.number_format = '@'
            if L == 'Y':
                c.fill = yel
    last = 5 + len(rows)
    if rows:
        rng = lambda L: f'{L}6:{L}{last}'  # noqa: E731
        ws.conditional_formatting.add(rng('M'), ColorScaleRule(start_type='min', start_color='F8696B', mid_type='percentile',
                                                               mid_value=50, mid_color='FFEB84', end_type='max', end_color='63BE7B'))
        ws.conditional_formatting.add(rng('R'), ColorScaleRule(start_type='min', start_color='F8696B', mid_type='percentile',
                                                               mid_value=50, mid_color='FFFFFF', end_type='max', end_color='63BE7B'))
    widths = {'A': 6, 'B': 20, 'C': 21, 'D': 9, 'E': 8, 'R': 13, 'Z': 14}
    for L in letters:
        ws.column_dimensions[L].width = widths.get(L, 11.5)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[3].height = 42
    ws.row_dimensions[5].height = 36
    ws.freeze_panes = 'F6'

    w2 = wb.create_sheet('校验')
    w2.append(['类别', '内容'])
    for k, v in meta['summary']:
        w2.append([k, v])
    w2.append([])
    w2.append(['提示', '（无）' if not WARN else f'共 {len(WARN)} 条，逐条看一眼'])
    for m in WARN:
        w2.append(['', m])
    w2.column_dimensions['A'].width = 18
    w2.column_dimensions['B'].width = 120
    for row in w2.iter_rows():
        for c in row:
            c.font = Font(name=F, size=10, bold=(c.row == 1))
    wb.save(out)


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description='KANS 越南 C-ads 品广素材投放决策汇总')
    ap.add_argument('inputs', nargs='+', help='上传的文件或文件夹（自动识别类型）')
    ap.add_argument('--fx', type=float, required=True, help='汇率：1 USD = ? VND（必须跟用户确认，例 26310）')
    ap.add_argument('--period', help='分析周期，例 "2026-09-23 ~ 2026-09-29"；不填就用 TTMS 报告里的周期')
    ap.add_argument('--status-date', help='状态/日预算的取数日期，例 10.2；不填就用 Campaign Report 文件名里的日期')
    ap.add_argument('--out', default='.', help='输出文件路径或目录（目录则自动命名为 …汇总_MMDD-MMDD.xlsx）')
    a = ap.parse_args()

    k = classify(a.inputs)
    ttms_rows = []
    for f in k['ttms_text']:
        for enc in ('utf-8-sig', 'utf-16', 'gb18030'):
            try:
                txt = open(f, encoding=enc).read()
                break
            except UnicodeError:
                continue
        ttms_rows += parse_ttms_text(txt, f)
    for f in k['ttms_table']:
        df = pd.read_csv(f, dtype=str, encoding='utf-8-sig') if f.lower().endswith('.csv') else pd.read_excel(f, dtype=str)
        ttms_rows += parse_ttms_table(df, f)
    if not ttms_rows:
        sys.exit('没读到任何 TTMS 报告数据。请先把报告页内容粘成 txt，或填 TTMS结案录入模板.csv')
    vids = ttms_by_video(ttms_rows)
    gm, gm_labels = load_gmvmax(k['gmvmax'], a.fx)
    orders = load_orders(k['orders'])
    camp, camp_date = load_campaign(k['campaign'])
    if not k['gmvmax']:
        warn('没有 GMV MAX 导出：N–R、T、U 全部留空，表按质量分看')
    if not k['orders']:
        warn('没有达人订单明细：V/W（付费/自然流占比）留空')
    if not k['campaign']:
        warn('没有 Campaign Report：KOL 名取 GMV MAX 的账号名，状态/现行日预算留空')

    rows = build_rows(vids, gm, orders, camp)
    low = [r for r in rows if r.get('n_orders') is not None and r['n_orders'] < 15]
    if low:
        warn(f'{len(low)} 个素材已结算订单不到 15 单，V/W 仅供参考：'
             + '、'.join(f"{r.get('B', '')}({r['n_orders']}单)" for r in low))
    tiny = [r for r in rows if isinstance(r.get('N'), float) and r['N'] < 10]
    if tiny:
        warn(f'{len(tiny)} 个素材 GMV MAX 消耗不到 $10，GMVMAX ROI / 获客成本没有参考意义（几乎没花钱，收入多为自然成交）：'
             + '、'.join(f"{r.get('B', '')}(${r['N']:.2f}, ROI {r['O']:.1f})" if r.get('O') is not None
                         else f"{r.get('B', '')}(${r['N']:.2f})" for r in tiny))

    cycles = Counter(c for d in vids.values() for c in d['cycles'])
    period = a.period or (cycles.most_common(1)[0][0] if cycles else '（未填）')
    if len(cycles) > 1:
        warn(f'报告的分析周期不一致：{dict(cycles)}，表头按「{period}」写')
    sd = a.status_date or camp_date or ''
    n_rep = sum(len(d['reports']) for d in vids.values())
    line2 = (f'分析周期 {period} ｜ 归因期限 7 天 ｜ 数据源：TTMS 种草广告报告({n_rep}份)'
             + (f" + GMV MAX Creative({'/'.join(gm_labels)})" if gm_labels else '')
             + (' + 达人订单明细' if k['orders'] else '') + f' ｜ 汇率 {a.fx:g}')
    line3 = ('口径：G=F/H；M=等权归一化 (mm(I)+mm(J)+(1-mm(G)))/3，批内可比、跨批不可比；'
             'K 优先取报告页 GMV，页面缺失或与 F×L 差>2% 时用 F×L 反算；R=GMVMAX收入/(品广消耗+GMVMAX消耗)，分子不含品广GMV；'
             'V=有店铺广告佣金率的订单GMV/该素材总GMV（按内容ID聚合，仅已结算）；T/U 在 GMVMAX 曝光<1000 时留空；'
             + (f'E/X 取 {sd} 当天状态；' if sd else '') + '单数<15 的 V/W 不具统计意义')
    summary = [('TTMS 报告', f'{n_rep} 份 → {len(vids)} 个素材'),
               ('GMV MAX 文件', '、'.join(os.path.basename(x) for x in k['gmvmax']) or '无'),
               ('达人订单', '、'.join(os.path.basename(x) for x in k['orders']) or '无'),
               ('Campaign Report', '、'.join(os.path.basename(x) for x in k['campaign']) or '无'),
               ('没认出来的文件', '、'.join(os.path.basename(x) for x in k['unknown']) or '无'),
               ('汇率', f'1 USD = {a.fx:g} VND')]
    camp_only = [v for v in camp if v not in vids and '投放中' in camp[v]['status']]
    if camp_only:
        def _start(v):
            m = re.match(r'\s*(\d{4})', camp[v]['groups'][0])
            return f"{m.group(1)[:2]}/{m.group(1)[2:]}起投" if m else ''
        warn('这些素材在 Campaign Report 里「投放中」但没有 TTMS 报告（起投太晚还没建、或漏传了报告）：'
             + '、'.join(f"{camp[v]['kol']}({v}，{_start(v)})" for v in camp_only))
    out = a.out
    if os.path.isdir(out) or out.endswith(('/', os.sep)) or not out.lower().endswith('.xlsx'):
        tag = '-'.join(re.sub(r'^\d{4}-(\d{2})-(\d{2})$', r'\1\2', x.strip()) for x in period.split('~')) if '~' in period else 'out'
        out = os.path.join(out, f'KANS_C-ads品广素材投放决策汇总_{tag}.xlsx')
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    write_xlsx(rows, out, {'line2': line2, 'line3': line3, 'summary': summary})

    print(f'✅ 已生成：{out}')
    print(f'   {len(rows)} 个素材 ｜ 周期 {period} ｜ 汇率 {a.fx:g}')
    for kk, v in summary:
        print(f'   {kk}：{v}')
    top = [r for r in rows if r.get('R') is not None][:3]
    if top:
        print('   总营销ROI 前三：' + '；'.join(f"{r['B']}({r['prod']}) {r['R']:.2f}" for r in top))
    best = sorted([r for r in rows if r.get('M') is not None], key=lambda r: -r['M'])[:3]
    if best:
        print('   品广质量分前三：' + '；'.join(f"{r['B']}({r['prod']}) {r['M']:.3f}" for r in best))
    print(f'⚠️ 校验提示 {len(WARN)} 条：' if WARN else '校验提示：无')
    for m in WARN:
        print('   - ' + m)


if __name__ == '__main__':
    main()
