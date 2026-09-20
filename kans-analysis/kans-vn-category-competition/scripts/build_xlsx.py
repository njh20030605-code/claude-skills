# -*- coding: utf-8 -*-
"""
build_xlsx.py —— 输出人民币口径两周对比工作簿（4 个 sheet）
  两周对比_CNY   : 店铺层，本周排名 / 上周排名 / 排名变化 / 中文名 / 七组「本周|上周」指标，KANS 行黄底
  渠道结构对比    : 大盘小结 + 分店铺渠道占比与 pp 变化 + 异动标注
  数据源_USD     : 平台原始美元值（主表公式源，改汇率全表重算）
  上周在榜本周掉出 : 掉榜清单
"""
import argparse
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import common as C

p = argparse.ArgumentParser()
p.add_argument("--cur", required=True)
p.add_argument("--pri", required=True)
p.add_argument("--rate", type=float, default=C.DEFAULT_RATE)
p.add_argument("--rate-note", default="")
p.add_argument("--cur-label", default="本周")
p.add_argument("--pri-label", default="上周")
p.add_argument("--out", default="两周对比_CNY.xlsx")
a = p.parse_args()

cur, pri = C.read_export(a.cur), C.read_export(a.pri)
B = C.build(cur, pri)
M = B["market"]
c1, c0 = cur.set_index("store"), pri.set_index("store")
RAW = ["gmv"] + C.CH + ["aov"]
GROUPS = ["总GMV\nGMV", "达人视频GMV\nAffiliate Video", "店播短视频GMV\nSelf-Account Video",
          "店铺自播GMV\nSelf-Account Live", "达人直播GMV\nAffiliate Live",
          "商城GMV\nShop Tab", "自播客单价\nSelf-Live AOV"]
F = "Arial"
thin = Side(style="thin", color="D9D9D9")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="404040")
SUB = PatternFill("solid", fgColor="DCE6F1")
GRY = PatternFill("solid", fgColor="F2F2F2")
YEL = PatternFill("solid", fgColor="FFF2A8")
ORG = PatternFill("solid", fgColor="FCE4D6")
wb = Workbook()

# ---------- 数据源_USD ----------
src = wb.create_sheet("数据源_USD")
src["A1"] = "平台原始导出（USD）。主表按此表 × 汇率换算，改汇率即全表重算。"
src["A1"].font = Font(name=F, size=9, color="808080")
src["A3"] = "Store Name"
for j, c in enumerate(RAW, start=2):
    src.cell(row=3, column=j, value="W1_" + c)
    src.cell(row=3, column=j + 7, value="W0_" + c)
for i, s in enumerate(cur.store):
    r = 4 + i
    src.cell(row=r, column=1, value=s)
    for j, c in enumerate(RAW, start=2):
        src.cell(row=r, column=j, value=float(c1.loc[s, c]))
        if s in c0.index:
            src.cell(row=r, column=j + 7, value=float(c0.loc[s, c]))
for c in src[3]:
    c.font = Font(name=F, size=9, bold=True)
src.column_dimensions["A"].width = 32

# ---------- 主表 ----------
ws = wb.create_sheet("两周对比_CNY", 0)
ws["A1"] = f"越南美妆个护 TikTok Shop 店铺 GMV 两周对比（人民币口径）"
ws["A1"].font = Font(name=F, size=13, bold=True)
ws["A2"] = f"{a.cur_label}（排名依据） vs {a.pri_label}　单位：元"
ws["A2"].font = Font(name=F, size=9, color="808080")
ws["A3"] = "汇率 USD→CNY："
ws["A3"].font = Font(name=F, size=9, bold=True)
ws["C3"] = a.rate
ws["C3"].font = Font(name=F, size=9, bold=True, color="0000FF")
ws["C3"].fill = PatternFill("solid", fgColor="FFFF00")
ws["C3"].number_format = "0.0000"
ws["D3"] = a.rate_note or "改此单元格全表自动重算。"
ws["D3"].font = Font(name=F, size=9, color="808080")

H1, H2, R0 = 5, 6, 7
fixed = [f"{a.cur_label}排名", f"{a.pri_label}排名", "排名变化", "店铺名\nStore Name", "品牌中文名"]
for j, hh in enumerate(fixed, start=1):
    ws.merge_cells(start_row=H1, start_column=j, end_row=H2, end_column=j)
    c = ws.cell(row=H1, column=j, value=hh)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
col = 6
for g in GROUPS:
    ws.merge_cells(start_row=H1, start_column=col, end_row=H1, end_column=col + 1)
    c = ws.cell(row=H1, column=col, value=g)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k, t in enumerate([a.cur_label, a.pri_label]):
        x = ws.cell(row=H2, column=col + k, value=t)
        x.fill = SUB if k == 0 else GRY
        x.font = Font(name=F, size=9, bold=True)
        x.alignment = Alignment(horizontal="center")
    col += 2

for i, s in enumerate(cur.store):
    r, sr = R0 + i, 4 + i
    inw0 = s in c0.index
    ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=2,
            value=int(pri.index[pri.store == s][0]) + 1 if inw0 else None
            ).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=3,
            value=f'=IF(B{r}="","新进",IF(B{r}-A{r}>0,"↑"&(B{r}-A{r}),'
                  f'IF(B{r}-A{r}<0,"↓"&(A{r}-B{r}),"—")))'
            ).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=4, value=s)
    ws.cell(row=r, column=5, value=C.label(s))
    col = 6
    for k in range(7):
        cc = ws.cell(row=r, column=col, value=f"=数据源_USD!{get_column_letter(2+k)}{sr}*$C$3")
        cc.number_format = "#,##0"
        c2 = ws.cell(row=r, column=col + 1)
        if inw0:
            c2.value = f"=数据源_USD!{get_column_letter(9+k)}{sr}*$C$3"
            c2.number_format = "#,##0"
        col += 2
    for j in range(1, 20):
        cell = ws.cell(row=r, column=j)
        cell.border = bd
        cell.font = Font(name=F, size=10)
        if j >= 7 and j % 2 == 1:
            cell.fill = GRY
    if s == C.KANS:
        for j in range(1, 20):
            ws.cell(row=r, column=j).fill = YEL
            ws.cell(row=r, column=j).font = Font(name=F, size=10, bold=True)

last = R0 + len(cur) - 1
ws.freeze_panes = f"F{R0}"
ws.auto_filter.ref = f"A{H2}:S{last}"
for j, w in enumerate([9, 9, 10, 30, 22] + [13, 13] * 7, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[H1].height = 34
notes = [
    "口径说明：",
    f"1. 按{a.cur_label} Top{len(cur)} 为基准，排名按总 GMV 降序。",
    f"2. 「新进」= 上周不在榜，不等于上周为 0；两周门槛不同（{cur.gmv.min():,.0f} vs {pri.gmv.min():,.0f} USD）。",
    "3. 各渠道 GMV 相加 ≠ 总 GMV（字段口径重叠），不可直接当结构占比。",
    "4. 客单价为平台均值字段，不做汇总；整数美元，±1 约 ±4-5%。",
    "5. 品牌中文名带 ? 的为归属未确认，需人工核对。",
]
for k, t in enumerate(notes):
    c = ws.cell(row=last + 3 + k, column=4, value=t)
    c.font = Font(name=F, size=9, bold=(k == 0), color="808080" if k else "000000")

# ---------- 渠道结构对比 ----------
w2 = wb.create_sheet("渠道结构对比", 1)
w2["A1"] = "渠道结构两周对比（占比与升降）"
w2["A1"].font = Font(name=F, size=13, bold=True)
w2["A2"] = ("占比分母 = 五渠道之和（非总GMV）。顺序与主表一致。")
w2["A2"].font = Font(name=F, size=9, color="808080")
w2["A4"] = f"① 大盘渠道结构（仅两周均在榜的 {M['n']} 家，口径可比）"
w2["A4"].font = Font(name=F, size=11, bold=True)
for j, hh in enumerate(["渠道", f"{a.cur_label}占比", f"{a.pri_label}占比", "变化(pp)",
                        f"{a.cur_label}金额(元)", f"{a.pri_label}金额(元)", "环比"], start=1):
    c = w2.cell(row=5, column=j, value=hh)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = Alignment(horizontal="center")
for k, nm in enumerate(C.NM):
    r = 6 + k
    w2.cell(row=r, column=1, value=nm)
    w2.cell(row=r, column=2, value=M["share1"][k] / 100).number_format = "0.0%"
    w2.cell(row=r, column=3, value=M["share0"][k] / 100).number_format = "0.0%"
    w2.cell(row=r, column=4, value=M["share1"][k] - M["share0"][k]).number_format = "+0.0;-0.0;0.0"
    w2.cell(row=r, column=5, value=M["ch1"][k] * a.rate).number_format = "#,##0"
    w2.cell(row=r, column=6, value=M["ch0"][k] * a.rate).number_format = "#,##0"
    w2.cell(row=r, column=7, value=M["ch_chg"][k] / 100).number_format = "+0.0%;-0.0%;0.0%"
w2.cell(row=11, column=1, value="五渠道合计").font = Font(name=F, size=10, bold=True)
w2.cell(row=11, column=5, value="=SUM(E6:E10)").number_format = "#,##0"
w2.cell(row=11, column=6, value="=SUM(F6:F10)").number_format = "#,##0"
w2.cell(row=11, column=7, value="=E11/F11-1").number_format = "+0.0%;-0.0%;0.0%"
w2.cell(row=12, column=1, value=f"总 GMV 口径环比 {M['gmv_chg']:+.1f}%（与品牌总GMV环比对比时用这个）"
        ).font = Font(name=F, size=9, color="808080")
for r in range(5, 12):
    for j in range(1, 8):
        w2.cell(row=r, column=j).border = bd

w2.cell(row=14, column=1, value="② 分店铺渠道结构").font = Font(name=F, size=11, bold=True)
G1, G2, D0 = 15, 16, 17
fx = [f"{a.cur_label}排名", "排名变化", "店铺名", "品牌中文名", f"总GMV\n{a.cur_label}(元)",
      f"总GMV\n{a.pri_label}(元)", "总GMV\n环比", f"渠道合计/总GMV\n{a.cur_label}",
      f"渠道合计/总GMV\n{a.pri_label}"]
for j, hh in enumerate(fx, start=1):
    w2.merge_cells(start_row=G1, start_column=j, end_row=G2, end_column=j)
    c = w2.cell(row=G1, column=j, value=hh)
    c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
col = 10
for nm in C.NM:
    w2.merge_cells(start_row=G1, start_column=col, end_row=G1, end_column=col + 2)
    c = w2.cell(row=G1, column=col, value=nm + "占比")
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = Alignment(horizontal="center", vertical="center")
    for k, t in enumerate([a.cur_label, a.pri_label, "变化pp"]):
        x = w2.cell(row=G2, column=col + k, value=t)
        x.fill = SUB if k == 0 else (GRY if k == 1 else PatternFill("solid", fgColor="EDEDED"))
        x.font = Font(name=F, size=9, bold=True)
        x.alignment = Alignment(horizontal="center")
    col += 3
w2.merge_cells(start_row=G1, start_column=25, end_row=G2, end_column=25)
c = w2.cell(row=G1, column=25, value="异动标注")
c.font = Font(name=F, size=9, bold=True, color="FFFFFF")
c.fill = HDR
c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

S = "数据源_USD"
for i, s in enumerate(cur.store):
    r, sr, mr = D0 + i, 4 + i, R0 + i
    inw0 = s in c0.index
    w2.cell(row=r, column=1, value=f"=两周对比_CNY!A{mr}").alignment = Alignment(horizontal="center")
    w2.cell(row=r, column=2, value=f"=两周对比_CNY!C{mr}").alignment = Alignment(horizontal="center")
    w2.cell(row=r, column=3, value=s)
    w2.cell(row=r, column=4, value=C.label(s))
    w2.cell(row=r, column=5, value=f"=两周对比_CNY!F{mr}").number_format = "#,##0"
    w2.cell(row=r, column=6, value=f'=IF(两周对比_CNY!G{mr}="","",两周对比_CNY!G{mr})'
            ).number_format = "#,##0"
    w2.cell(row=r, column=7, value=f'=IF(F{r}="","",E{r}/F{r}-1)'
            ).number_format = "+0.0%;-0.0%;0.0%"
    w2.cell(row=r, column=8, value=f"=SUM({S}!C{sr}:G{sr})/{S}!B{sr}").number_format = "0.0%"
    w2.cell(row=r, column=9, value=f'=IF({S}!I{sr}="","",SUM({S}!J{sr}:N{sr})/{S}!I{sr})'
            ).number_format = "0.0%"
    col = 10
    for k in range(5):
        a1, a0 = get_column_letter(3 + k), get_column_letter(10 + k)
        w2.cell(row=r, column=col,
                value=f'=IFERROR({S}!{a1}{sr}/SUM({S}!$C{sr}:$G{sr}),"")').number_format = "0.0%"
        w2.cell(row=r, column=col + 1,
                value=f'=IF({S}!I{sr}="","",IFERROR({S}!{a0}{sr}/SUM({S}!$J{sr}:$N{sr}),""))'
                ).number_format = "0.0%"
        L1, L2 = get_column_letter(col), get_column_letter(col + 1)
        w2.cell(row=r, column=col + 2,
                value=f'=IF(OR({L1}{r}="",{L2}{r}=""),"",({L1}{r}-{L2}{r})*100)'
                ).number_format = "+0.0;-0.0;0.0"
        col += 3
    flags = []
    if inw0:
        chg = c1.loc[s, "gmv"] / c0.loc[s, "gmv"] - 1
        if c0.loc[s, "gmv"] >= 20000 and chg >= 1.0:
            flags.append(f"GMV +{chg*100:.0f}%")
        if chg <= -0.3:
            flags.append(f"GMV {chg*100:.0f}%")
        v1 = np.array([c1.loc[s, c] for c in C.CH], float)
        v0 = np.array([c0.loc[s, c] for c in C.CH], float)
        if v0.sum() > 0 and c1.loc[s, "gmv"] >= 100000:
            for j, nm in enumerate(C.NM):
                pp = (v1[j] / v1.sum() - v0[j] / v0.sum()) * 100
                if abs(pp) >= 20:
                    flags.append(f"{nm} {pp:+.0f}pp")
    else:
        flags.append(f"{a.cur_label}新进榜")
    txt = "；".join(flags)
    fc = w2.cell(row=r, column=25, value=txt)
    fc.font = Font(name=F, size=9)
    if txt and "新进" not in txt:
        fc.fill = ORG
    for j in range(1, 26):
        cell = w2.cell(row=r, column=j)
        cell.border = bd
        if j != 25:
            cell.font = Font(name=F, size=10)
        if j in (11, 14, 17, 20, 23):
            cell.fill = GRY
    if s == C.KANS:
        for j in range(1, 26):
            w2.cell(row=r, column=j).fill = YEL
            w2.cell(row=r, column=j).font = Font(name=F, size=10, bold=True)
w2.freeze_panes = f"E{D0}"
w2.auto_filter.ref = f"A{G2}:Y{D0+len(cur)-1}"
for j, w in enumerate([8, 9, 28, 22, 13, 13, 10, 14, 14] + [9, 9, 9] * 5 + [34], start=1):
    w2.column_dimensions[get_column_letter(j)].width = w
w2.row_dimensions[G1].height = 32

# ---------- 掉榜 ----------
out = [s for s in pri.store if s not in set(cur.store)]
w3 = wb.create_sheet("上周在榜本周掉出")
w3["A1"] = f"{a.pri_label}在榜、{a.cur_label}未在榜的店铺，共 {len(out)} 家（单位：元）"
w3["A1"].font = Font(name=F, size=11, bold=True)
for j, hh in enumerate([f"{a.pri_label}排名", "店铺名", "品牌中文名", f"{a.pri_label}总GMV(元)"],
                       start=1):
    c = w3.cell(row=3, column=j, value=hh)
    c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    c.fill = HDR
    c.alignment = Alignment(horizontal="center")
for i, s in enumerate(out):
    r = 4 + i
    w3.cell(row=r, column=1, value=int(pri.index[pri.store == s][0]) + 1
            ).alignment = Alignment(horizontal="center")
    w3.cell(row=r, column=2, value=s)
    w3.cell(row=r, column=3, value=C.label(s))
    w3.cell(row=r, column=4, value=float(c0.loc[s, "gmv"]) * a.rate).number_format = "#,##0"
    for j in range(1, 5):
        w3.cell(row=r, column=j).border = bd
        w3.cell(row=r, column=j).font = Font(name=F, size=10)
for j, w in enumerate([10, 32, 22, 16], start=1):
    w3.column_dimensions[get_column_letter(j)].width = w

del wb["Sheet"]
wb.save(a.out)
print("→", a.out, "（含公式，务必再跑 recalc.py）")
