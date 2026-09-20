"""
common.py — 三张图共用的工具：业务参数、中文字体、配色、数字格式、数据读取/解析。

数据读取支持两种来源：
  1) 干净 CSV（推荐，见 examples/）
  2) 平台原始导出 xlsx（直播间 Campaign 导出 / 大盘类目导出），自动解析

所有脚本都从这里 import，保证三张图的字体、配色、口径一致。
"""
import os
import re
import glob
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm

# ======================================================================
# 业务参数（换成你自己的；三张图的标题 / 图例 / 默认值都从这里取）
# ======================================================================
# 示例数据：换成你自己的导出
PARAMS = dict(
    brand="你的品牌",              # 图例里的品牌简称，例：KANS
    shop_label="你的品牌 直播间",   # 图1 标题用，例：KANS 越南直播间
    market="示例市场",              # 例：越南
    category="示例类目",            # 大盘类目名，例：Serums & Essences
    platform="TikTok Shop",         # 大盘所在平台
    market_source="类目大盘导出",    # 大盘数据来源名，例：TTMS
    self_source="直播间 Campaign 导出",  # 自播数据来源名
    year="2026",                    # 标题里的年份（date 列只有 MM-DD）
    usd_rate=6.8,                   # 1 USD = ? RMB（图1 默认汇率，可用 --rate 覆盖）
    breakeven=4.1,                  # 退款后盈亏线 ROI（图1 默认，可用 --breakeven 覆盖）
    promo_dates=["06-06", "06-18"], # 平台大促日（图2 默认，可用 --promo 覆盖；示例数据里的日期）
    ramp_dates=["06-17", "06-19"],  # 自播拉量日（图2 默认，可用 --ramp 覆盖；示例数据里的日期）
)

# 平台原始 xlsx 的表头关键字（后台是英文界面时改这里；直接喂 CSV 可忽略）
XLSX_HEADERS = dict(
    self_time="启动时间", self_spend="成本", self_gmv="总收入",          # 直播间 Campaign 导出
    mkt_date="日期", mkt_short="短视频", mkt_live="直播", mkt_card="商品卡",  # 大盘类目导出
)

# ----------------------------------------------------------------------
# 中文字体：按候选列表顺序找第一个存在的；可用环境变量 CJK_FONT 指定字体文件路径。
# 找不到时退回 matplotlib 默认字体（中文会显示成方块，但不报错）。
# ----------------------------------------------------------------------
_CJK_CANDIDATES = [
    os.environ.get("CJK_FONT", ""),
    # macOS
    "/System/Library/Fonts/PingFang.ttc",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
    "/Library/Fonts/Arial Unicode.ttf",
    # Linux
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Black.ttc",
    "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    # Windows
    "C:/Windows/Fonts/msyh.ttc",
    "C:/Windows/Fonts/simhei.ttf",
]

def setup_cjk_font():
    """注册并设置中文字体，返回 (regular_name, BOLD_FontProperties)。找不到 CJK 字体时不报错。"""
    reg_path = None
    blk_path = None
    for p in _CJK_CANDIDATES:
        if p and os.path.exists(p):
            try:
                fm.fontManager.addfont(p)
            except Exception:
                continue
            if reg_path is None:
                reg_path = p
            if "Black" in p or "Bold" in p:
                blk_path = p
    if reg_path is None:
        # 最后兜底：在系统里搜任意带 CJK 的字体
        for f in fm.findSystemFonts():
            if re.search(r"(cjk|noto.*sc|wqy|hei|song|pingfang|hiragino)", f, re.I):
                try:
                    fm.fontManager.addfont(f); reg_path = f; break
                except Exception:
                    continue
    if reg_path is None:
        print("[warn] 未找到中文字体，退回默认字体（中文可能显示为方块）。可设置环境变量 CJK_FONT=/path/to/font.ttc")
        plt.rcParams["axes.unicode_minus"] = False
        return plt.rcParams["font.family"][0], fm.FontProperties(weight="bold")
    if blk_path is None:
        blk_path = reg_path
    plt.rcParams["font.family"] = fm.FontProperties(fname=reg_path).get_name()
    plt.rcParams["axes.unicode_minus"] = False
    blk = fm.FontProperties(fname=blk_path)
    if blk_path == reg_path:
        blk.set_weight("bold")
    return fm.FontProperties(fname=reg_path).get_name(), blk

# ----------------------------------------------------------------------
# 统一配色
# ----------------------------------------------------------------------
PALETTE = {
    # 图1 周对比
    "week_cycle": ["#4A90D9", "#16A085", "#8E5BD0", "#E0772B"],  # 周1/周2/周3...
    "roi":        "#D9572B",
    "breakeven":  "#888888",
    # 图2 相关性
    "market":     "#2C7FB8",
    "kans":       "#C0392B",   # 自有品牌线（键名保留，供其他脚本引用）
    "ramp":       "#1A8A78",
    "fit_full":   "#23303D",
    "fit_grey":   "#8B96A3",
    "promo_band": "#FBEFC6",
    # 图3 大盘堆叠
    "live":       "#2D6B8F",
    "short":      "#E0A23C",
    "card":       "#C9D2DA",
    "total_line": "#23303D",
    "ann_blue":   "#2D6B8F",
    "ann_red":    "#C0392B",
    "ann_teal":   "#1A8A78",
    "ann_grey":   "#7C8896",
    "grid":       "#ECECEC",
}

# ----------------------------------------------------------------------
# 数字格式
# ----------------------------------------------------------------------
def fmt_usd(v):
    """美元金额：>=1000(千) 显示 $X.XXM，否则 $XXXK。入参单位为『千USD』。"""
    return f"${v/1000:.2f}M" if v >= 1000 else f"${v:.0f}K"

def fmt_rmb(v):
    return f"¥{v:,.0f}"

# ----------------------------------------------------------------------
# 数据读取
# ----------------------------------------------------------------------
def _read_table(path):
    import pandas as pd
    if path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        return None  # 交给专用解析器
    return pd.read_csv(path)

def parse_kans_campaign_xlsx(path, start=None, end=None):
    """
    解析直播间 Campaign 导出 xlsx（按表头定位列，列序变动也能用）。
    需要的表头见 XLSX_HEADERS：self_time / self_spend / self_gmv（默认 '启动时间'、'成本'、'总收入'）。
    返回 DataFrame[date, spend_usd, gmv_usd]，按日聚合，已过滤无数据的排期占位行。
    """
    import openpyxl, pandas as pd
    from collections import defaultdict
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    def col(name):
        for i, h in enumerate(hdr):
            if h and name in str(h):
                return i + 1
        raise KeyError(f"列未找到: {name}")
    H = XLSX_HEADERS
    c_t, c_sp, c_gm = col(H["self_time"]), col(H["self_spend"]), col(H["self_gmv"])
    sp, gm = defaultdict(float), defaultdict(float)
    for r in range(2, ws.max_row + 1):
        t = ws.cell(r, c_t).value
        if not t:
            continue
        d = str(t)[:10]
        try:
            s = float(ws.cell(r, c_sp).value); g = float(ws.cell(r, c_gm).value)
        except (TypeError, ValueError):
            continue
        sp[d] += s; gm[d] += g
    rows = [(d, sp[d], gm[d]) for d in sorted(sp) if (sp[d] > 0 or gm[d] > 0)]
    df = pd.DataFrame(rows, columns=["date", "spend_usd", "gmv_usd"])
    df["date"] = df["date"].str[5:]  # 'YYYY-MM-DD' -> 'MM-DD'
    return _slice(df, start, end)

def parse_ttms_serums_xlsx(path, start=None, end=None):
    """
    解析大盘类目导出 xlsx（例：TTMS 类目日度导出）。
    需要的表头见 XLSX_HEADERS：mkt_date / mkt_short / mkt_live / mkt_card
    （默认 '日期'、'短视频GMV'、'直播GMV'、'商品卡GMV'，单位千USD）。
    返回 DataFrame[date, short_k, live_k, card_k]，自动跳过『区间最大值』等汇总行。
    """
    import openpyxl, pandas as pd
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    def col(*keys):
        for i, h in enumerate(hdr):
            if h and all(k in str(h) for k in keys):
                return i + 1
        raise KeyError(f"列未找到: {keys}")
    H = XLSX_HEADERS
    c_d = col(H["mkt_date"]); c_s = col(H["mkt_short"]); c_l = col(H["mkt_live"]); c_c = col(H["mkt_card"])
    rows = []
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, c_d).value
        if not d or not re.match(r"^\d{1,2}-\d{1,2}", str(d).strip()):
            continue  # 跳过『区间最大值』等汇总行
        rows.append((str(d).strip(),
                     float(ws.cell(r, c_s).value),
                     float(ws.cell(r, c_l).value),
                     float(ws.cell(r, c_c).value)))
    df = pd.DataFrame(rows, columns=["date", "short_k", "live_k", "card_k"])
    return _slice(df, start, end)

def _slice(df, start, end):
    if start:
        df = df[df["date"] >= start]
    if end:
        df = df[df["date"] <= end]
    return df.sort_values("date").reset_index(drop=True)

def load_kans_daily(path, start=None, end=None):
    """读自播日度数据。.csv 需含 date,spend_usd,gmv_usd；.xlsx 走 Campaign 解析器。"""
    import pandas as pd
    df = _read_table(path)
    if df is None:
        return parse_kans_campaign_xlsx(path, start, end)
    return _slice(df, start, end)

def load_market_daily(path, start=None, end=None):
    """读大盘日度数据。.csv 需含 date,short_k,live_k,card_k；.xlsx 走类目导出解析器。"""
    import pandas as pd
    df = _read_table(path)
    if df is None:
        return parse_ttms_serums_xlsx(path, start, end)
    return _slice(df, start, end)

def style_axes(ax, hide=("top", "right")):
    for s in hide:
        ax.spines[s].set_visible(False)
