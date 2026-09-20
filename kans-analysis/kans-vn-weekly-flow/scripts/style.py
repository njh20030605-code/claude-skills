# -*- coding: utf-8 -*-
"""共用 openpyxl / matplotlib 样式（对齐 KANS 周报既有口径）"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RED = 'C00000'        # 表头红底
BLUE = '2E75B6'       # 本期列蓝表头
DARKRED = 'C00000'    # 环比 ↑
GREEN = '00B050'      # 环比 ↓
YELLOW = 'FFE699'     # 高亮黄
DEEPYELLOW = 'FFC232'
GREY = 'F2F2F2'
NOTEBLUE = 'DDEBF7'

THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def hdr(ws, row, ncol, fill=RED, start=1):
    for c in range(start, start + ncol):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill('solid', fgColor=fill)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


def title(ws, row, text, ncol, fill=RED):
    ws.cell(row=row, column=1, value=text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row=row, column=1)
    c.fill = PatternFill('solid', fgColor=fill)
    c.font = Font(bold=True, color='FFFFFF', size=12)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 22


def widths(ws, spec):
    for i, w in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def pct_color(cell, v, invert=False):
    """红=↑ 绿=↓（中国财务口径）；invert=True 时语义相反（用于流量精度表）"""
    if v is None:
        return
    up, dn = (DARKRED, GREEN) if not invert else (GREEN, DARKRED)
    cell.font = Font(color=up if v > 0 else dn, bold=True, size=10)


def note_block(ws, row, lines, ncol):
    for i, t in enumerate(lines):
        r = row + i
        ws.cell(row=r, column=1, value=t)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        c = ws.cell(row=r, column=1)
        c.fill = PatternFill('solid', fgColor=NOTEBLUE)
        c.font = Font(size=9, color='1F4E79')
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    return row + len(lines)


# ---------------- matplotlib ----------------
def mpl_setup():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.sans-serif'] = ['Noto Sans CJK JP', 'WenQuanYi Zen Hei', 'DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    plt.rcParams['figure.dpi'] = 130
    plt.rcParams['savefig.dpi'] = 160
    plt.rcParams['axes.edgecolor'] = '#8C8C8C'
    plt.rcParams['axes.linewidth'] = 0.8
    plt.rcParams['grid.color'] = '#E0E0E0'
    return plt


C_PREV = '#9FB6CD'    # 上期 灰蓝
C_CUR = '#C0392B'     # 本期 红
C_ROI_PREV = '#7F8C8D'
C_ROI_CUR = '#E67E22'
C_LIVE = '#C0392B'
C_VIDEO = '#2E86C1'
C_OTHER = '#95A5A6'
