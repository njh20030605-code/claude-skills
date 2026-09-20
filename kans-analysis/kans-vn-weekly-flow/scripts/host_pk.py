#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
主播 PK 表 + 时段汇总 (KANS VN SKINCARE 直播间周报 Step 11)

输入：REPORT PERFORMANCE _ KANS ... SKINCARE.csv（平台原始导出）
关键：直接用 per-主播原生列（不要用 GMV Tổng 累计列做差值还原）：
  col13 GMV Tổng(当日累计) | col14 GMV Host(本段主播GMV) | col15 Spend Tổng(累计)
  col16 Spend Host | col17 ROI | col18 GMV/H | col19 Views
  col21 CTR | col23 Tap-through | col25/26 CTOR
该表偶有坏格——脚本自动清洗：
  · GMV Host<=0 或被错填成"当日累计"(=GMV Tổng 且非首段) → 用 GMV Tổng 累计差值(按时长拆相邻坏段)还原
  · Spend Host<=0 → 同法用 Spend Tổng 还原
  · CTR/CTOR > 60% → 视为坏值剔除(直播 CTR 正常 15-40%)

口径：直播间 GMV = VND/3800 → RMB；单h GMV = GMV(元)/时长；ROI = GMV/消耗。
班次(5)：早06-10 / 午10-14 / 下午14-18 / 晚18-22 / 夜22-02(含00:00-05:59跨夜段)。

用法：
  python host_pk.py --csv "REPORT PERFORMANCE ... SKINCARE.csv" \
      --cur 22/06 --prev 15/06 --year 2026 \
      --out "KANS_VN_主播PK+时段汇总.xlsx"
--cur / --prev 给该周【周一】日期(DD/MM)，自动展开 7 天(周一→周日)。
"""
import csv, argparse, collections, datetime, json, os

def vnd(x):
    s = str(x).replace('.', '').replace(',', '').strip()
    try: return float(s)
    except: return 0.0

def pct(x):
    s = str(x).replace('.', '').replace(',', '.').replace('%', '').strip()
    try: return float(s)
    except: return None

def t2m(t):
    t = str(t).strip()
    return int(t.split(':')[0]) * 60 + int(t.split(':')[1]) if ':' in t else None


def dur_h(dcol, start, end):
    v=str(dcol).strip().replace(',', '.')
    if v:
        try:
            f=float(v)
            if f>0: return f
        except: pass
    a=t2m(start); b=t2m(end)
    if a is None or b is None: return 0.0
    if b<=a: b+=24*60
    return (b-a)/60.0

def shift(st):
    if st < 360:  return '夜班 22-02'   # 00:00-05:59 跨夜尾
    if st < 600:  return '早班 06-10'
    if st < 840:  return '午班 10-14'
    if st < 1080: return '下午 14-18'
    if st < 1320: return '晚班 18-22'
    return '夜班 22-02'

ORDER = ['早班 06-10', '午班 10-14', '下午 14-18', '晚班 18-22', '夜班 22-02']
V = 3800.0

def week_days(monday_ddmm, year):
    d, m = map(int, monday_ddmm.split('/'))
    d0 = datetime.date(year, m, d)
    return ['%02d/%02d/%d' % ((d0 + datetime.timedelta(i)).day,
                              (d0 + datetime.timedelta(i)).month, year) for i in range(7)]

def load(csv_path):
    return list(csv.reader(open(csv_path, encoding='utf-8-sig')))

def extract(rows, days):
    """返回清洗后的 segment 列表（按日做累计差值修复坏格）。"""
    dayset = set(days)
    byday = collections.defaultdict(list)
    for x in rows[1:]:
        if len(x) < 27 or x[0] not in dayset or not x[3] or '#N/A' in str(x[9]):
            continue
        st = t2m(x[3])
        if st is None:
            continue
        byday[x[0]].append(dict(
            st=st, sh=shift(st), h=dur_h(x[5], x[3], x[4]), host=x[9],
            g_host=vnd(x[14]), g_cum=vnd(x[13]),
            sp_host=vnd(x[16]), sp_cum=vnd(x[15]),
            views=vnd(x[19]), ctr=pct(x[21]), ctor=pct(x[25]) or pct(x[26])))
    out = []
    for d, L in byday.items():
        L.sort(key=lambda z: z['st'])
        # --- 用累计列重建"干净的"每段值，作为坏格回退 ---
        def rebuild(cum_key):
            res = [None] * len(L); last = 0.0; pend = []
            for i, s in enumerate(L):
                c = s[cum_key]
                if c <= 0:            # 累计也缺 → 挂起
                    pend.append(i); continue
                if c < last:          # 跨夜重置
                    for j in pend: res[j] = 0.0
                    pend = []; res[i] = c; last = c; continue
                grp = pend + [i]; inc = c - last
                th = sum(L[j]['h'] for j in grp) or 1
                for j in grp: res[j] = inc * L[j]['h'] / th
                pend = []; last = c
            for j in pend: res[j] = 0.0
            return res
        gfix = rebuild('g_cum'); sfix = rebuild('sp_cum')
        for i, s in enumerate(L):
            g = s['g_host']
            # 坏格判定：<=0、或 ==当日累计(被错填)、或 显著偏离重建值
            bad = (g <= 0) or (i > 0 and abs(g - s['g_cum']) < 1) or \
                  (gfix[i] and abs(g - gfix[i]) > max(gfix[i], 1) * 0.6)
            s['g'] = gfix[i] if (bad and gfix[i]) else g
            sp = s['sp_host']
            spbad = (sp <= 0) or (i > 0 and abs(sp - s['sp_cum']) < 1) or \
                    (sfix[i] and abs(sp - sfix[i]) > max(sfix[i], 1) * 0.6)
            s['sp'] = (sfix[i] if (spbad and sfix[i]) else sp) or None
            if s['ctr'] is not None and s['ctr'] > 60:  s['ctr'] = None
            if s['ctor'] is not None and s['ctor'] > 60: s['ctor'] = None
            if s['g'] > 0:
                out.append(s)
    return out

def shift_totals(segs):
    s = collections.defaultdict(lambda: [0.0, 0.0, 0.0])  # gmv, cost, hours
    for x in segs:
        s[x['sh']][0] += x['g']; s[x['sh']][1] += (x['sp'] or 0); s[x['sh']][2] += x['h']
    return s

def pk_by_shift(segs):
    res = {}
    for k in ORDER:
        seg = [x for x in segs if x['sh'] == k]
        if not seg:
            continue
        hh = collections.defaultdict(lambda: [0, 0, 0, 0, [], [], 0])
        for x in seg:
            a = hh[x['host']]
            a[0] += x['g']; a[1] += (x['sp'] or 0); a[2] += x['h']; a[3] += x['views']; a[6] += 1
            if x['ctr'] is not None:  a[4].append(x['ctr'])
            if x['ctor'] is not None: a[5].append(x['ctor'])
        rows, tg, tsp, th, tv, tn, ac, aco = [], 0, 0, 0, 0, 0, [], []
        for host in sorted(hh, key=lambda z: -hh[z][0]):
            g, sp, h, v, ctr, ctor, n = hh[host]
            tg += g; tsp += sp; th += h; tv += v; tn += n; ac += ctr; aco += ctor
            rows.append([host, n, round(h), round(g / V / 1e4, 2),
                         round(sp / V / 1e4, 2) if sp else 0,
                         round(g / sp, 2) if sp else 0, round(g / V / (h or 1)),
                         round(v / n) if n else 0, round(g / V / v, 2) if v else 0,
                         round(sum(ctr) / len(ctr), 1) if ctr else None,
                         round(sum(ctor) / len(ctor), 1) if ctor else None])
        tot = ['合计', tn, round(th), round(tg / V / 1e4, 2),
               round(tsp / V / 1e4, 2) if tsp else 0, round(tg / tsp, 2) if tsp else 0,
               round(tg / V / th), round(tv / tn) if tn else 0,
               round(tg / V / tv, 2) if tv else 0,
               round(sum(ac) / len(ac), 1) if ac else None,
               round(sum(aco) / len(aco), 1) if aco else None]
        res[k] = {'rows': rows, 'tot': tot}
    return res

def summary(cur, prev):
    Sc, Sp = shift_totals(cur), shift_totals(prev)
    out = []; Tc = [0, 0, 0]; Tp = [0, 0, 0]
    for k in ORDER:
        gc, cc, hc = Sc.get(k, [0, 0, 0]); gp, cp, hp = Sp.get(k, [0, 0, 0])
        uc = gc / V / hc if hc else 0; up = gp / V / hp if hp else 0
        Tc = [Tc[0] + gc, Tc[1] + cc, Tc[2] + hc]; Tp = [Tp[0] + gp, Tp[1] + cp, Tp[2] + hp]
        out.append([k.split()[0], k.split()[1].replace('-', '–'),
                    round(hp), round(up), round(gp / cp, 2) if cp else 0,
                    round(hc), round(uc), round(gc / cc, 2) if cc else 0,
                    round((uc / up - 1) * 100) if up else 0])
    uc = Tc[0] / V / Tc[2] if Tc[2] else 0; up = Tp[0] / V / Tp[2] if Tp[2] else 0
    out.append(['合计/加权', '—', round(Tp[2]), round(up), round(Tp[0] / Tp[1], 2) if Tp[1] else 0,
                round(Tc[2]), round(uc), round(Tc[0] / Tc[1], 2) if Tc[1] else 0,
                round((uc / up - 1) * 100) if up else 0])
    return out

def build_xlsx(pk, summ, out_path, cur_label, prev_label):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    RED, GRN, TEAL = 'C0392B', '1E7E34', '2E5A66'
    bd = Border(*[Side('thin', color='D9D9D9')] * 4)
    emoji = {'早班 06-10': '🌅 早班 06–10', '午班 10-14': '☀️ 午班 10–14',
             '下午 14-18': '🌆 下午班 14–18', '晚班 18-22': '🌃 晚班 18–22',
             '夜班 22-02': '🌙 夜班 22–02'}
    def hdr(ws, cells, fill=RED):
        for c in cells:
            c.font = Font(bold=True, color='FFFFFF', size=10)
            c.fill = PatternFill('solid', fgColor=fill)
            c.alignment = Alignment('center', 'center', wrap_text=True); c.border = bd
    wb = openpyxl.Workbook()
    cols = ['主播', '场次', '时长h', 'GMV(万)', '消耗(万)', 'ROI', '单h GMV',
            '场均Views', 'UV价值', 'CTR', 'CTOR']
    ws = wb.active; ws.title = '主播PK表'
    ws.merge_cells('A1:K1')
    ws['A1'] = f'主播 PK 表（本周期 {cur_label}，5 班次）'
    ws['A1'].font = Font(bold=True, size=13, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=TEAL)
    ws['A1'].alignment = Alignment('center', 'center'); ws.row_dimensions[1].height = 24
    rix = 2
    for k in ORDER:
        if k not in pk: continue
        tot = pk[k]['tot']
        ws.merge_cells(start_row=rix, start_column=1, end_row=rix, end_column=11)
        c = ws.cell(rix, 1, f'{emoji[k]}　单h {tot[6]} ｜ ROI {tot[5]}')
        c.font = Font(bold=True, size=11, color='1F3A40')
        c.fill = PatternFill('solid', fgColor='DDEBEF')
        c.alignment = Alignment('left', 'center'); rix += 1
        for j, h in enumerate(cols, 1): ws.cell(rix, j, h)
        hdr(ws, [ws.cell(rix, j) for j in range(1, 12)]); rix += 1
        for row in pk[k]['rows']:
            for j, val in enumerate(row, 1):
                cc = ws.cell(rix, j, val if val is not None else '—')
                cc.border = bd; cc.alignment = Alignment('center'); cc.font = Font(size=10)
            ws.cell(rix, 1).font = Font(size=10, bold=True)
            if row[5] and row[5] < tot[5] - 0.01:
                ws.cell(rix, 6, f'▼{row[5]}'); ws.cell(rix, 6).font = Font(size=10, color=RED, bold=True)
            if row[6] and row[6] < tot[6]:
                ws.cell(rix, 7).font = Font(size=10, color=RED)
            rix += 1
        for j, val in enumerate(tot, 1):
            cc = ws.cell(rix, j, val if val is not None else '—')
            cc.border = bd; cc.alignment = Alignment('center')
            cc.font = Font(size=10, bold=True); cc.fill = PatternFill('solid', fgColor='F2F2F2')
        rix += 1
        ws.merge_cells(start_row=rix, start_column=1, end_row=rix, end_column=11)
        c = ws.cell(rix, 1, '💬 （在此补主播点评：谁稳/谁退/可加排黄金班/单人依赖风险等）')
        c.font = Font(size=9.5, color='8A6D00'); c.fill = PatternFill('solid', fgColor='FFF8E1')
        c.alignment = Alignment('left', 'center', wrap_text=True)
        ws.row_dimensions[rix].height = 26; rix += 2
    for col, w in zip('ABCDEFGHIJK', [12, 6, 7, 9, 9, 8, 9, 11, 8, 8, 8]):
        ws.column_dimensions[col].width = w
    # 时段汇总
    w2 = wb.create_sheet('时段汇总')
    w2.merge_cells('A1:I1')
    w2['A1'] = f'时段汇总（5 班次 ｜ 本期 {cur_label} vs 上期 {prev_label}）'
    w2['A1'].font = Font(bold=True, size=12, color='FFFFFF')
    w2['A1'].fill = PatternFill('solid', fgColor=TEAL); w2['A1'].alignment = Alignment('center', 'center')
    H = ['时段', '时间', '上期时长h', '上期单h GMV', '上期ROI', '本期时长h', '本期单h GMV', '本期ROI', '单h GMV变化']
    for j, x in enumerate(H, 1): w2.cell(3, j, x)
    hdr(w2, [w2.cell(3, j) for j in range(1, 10)])
    for row in summ:
        r = w2.max_row + 1
        vals = row[:8] + [f'{row[8]:+d}%']
        for j, val in enumerate(vals, 1):
            cc = w2.cell(r, j, val); cc.border = bd
            cc.alignment = Alignment('center'); cc.font = Font(size=10)
        if '合计' in str(row[0]):
            for j in range(1, 10):
                w2.cell(r, j).font = Font(size=10, bold=True)
                w2.cell(r, j).fill = PatternFill('solid', fgColor='F2F2F2')
        ch = row[8]
        w2.cell(r, 9).font = Font(size=10, bold=True, color=(GRN if ch > 0 else RED if ch < 0 else '555555'))
    for col, w in zip('ABCDEFGHI', [11, 9, 11, 13, 10, 11, 13, 10, 13]):
        w2.column_dimensions[col].width = w
    wb.save(out_path)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--csv', required=True)
    ap.add_argument('--cur', required=True, help='本周期周一 DD/MM')
    ap.add_argument('--prev', required=True, help='上周期周一 DD/MM')
    ap.add_argument('--year', type=int, default=datetime.date.today().year)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    rows = load(a.csv)
    cd, pd_ = week_days(a.cur, a.year), week_days(a.prev, a.year)
    cur, prev = extract(rows, cd), extract(rows, pd_)
    pk = pk_by_shift(cur); summ = summary(cur, prev)
    cl = f'{cd[0][:5]}–{cd[-1][:5]}'.replace('/', '.')
    pl = f'{pd_[0][:5]}–{pd_[-1][:5]}'.replace('/', '.')
    build_xlsx(pk, summ, a.out, cl, pl)
    print(json.dumps({'pk': pk, 'summary': summ}, ensure_ascii=False, indent=1))
    print('\nSaved ->', a.out)

if __name__ == '__main__':
    main()
